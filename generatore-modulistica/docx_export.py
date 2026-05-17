"""Esportazione della modulistica generata in un file Word (.docx)."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

_GRIGIO = RGBColor(0x55, 0x55, 0x55)


def _titolo_sezione(documento, testo):
    p = documento.add_paragraph()
    run = p.add_run(testo.upper())
    run.bold = True
    run.font.size = Pt(12)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)


def _campo_compilabile(documento, etichetta):
    p = documento.add_paragraph()
    run = p.add_run(f"{etichetta}: ")
    run.bold = True
    p.add_run("_" * 45)
    p.paragraph_format.space_after = Pt(4)


def crea_docx(dati, percorso_output):
    """Costruisce il documento Word a partire dai dati generati.

    `dati` deve essere conforme allo schema prodotto da generatore.py.
    """
    documento = Document()

    stile = documento.styles["Normal"]
    stile.font.name = "Calibri"
    stile.font.size = Pt(11)

    # Intestazione ente
    for riga in dati.get("modulo", {}).get("intestazione", []):
        p = documento.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(riga)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = _GRIGIO
        p.paragraph_format.space_after = Pt(0)

    # Titolo documento
    titolo = documento.add_paragraph()
    titolo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titolo.add_run(dati.get("tipologia", "Modulistica"))
    run.bold = True
    run.font.size = Pt(15)
    titolo.paragraph_format.space_before = Pt(10)
    titolo.paragraph_format.space_after = Pt(4)

    # Avviso documento didattico
    avviso = documento.add_paragraph()
    avviso.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = avviso.add_run(
        "Modello generato a scopo dimostrativo - da verificare e adattare "
        "alla modulistica ufficiale dell'ente"
    )
    run.italic = True
    run.font.size = Pt(8)
    run.font.color.rgb = _GRIGIO

    # Riferimenti di modulistica
    riferimenti = dati.get("riferimenti_modulistica", [])
    if riferimenti:
        _titolo_sezione(documento, "Riferimenti di modulistica")
        for rif in riferimenti:
            documento.add_paragraph(rif, style="List Bullet")

    # Checklist operativa
    checklist = dati.get("checklist", [])
    if checklist:
        _titolo_sezione(documento, "Checklist operativa")
        tabella = documento.add_table(rows=1, cols=5)
        tabella.style = "Light Grid Accent 1"
        intestazioni = ["Fatto", "Passo", "Responsabile", "Documenti", "Tempi / Note"]
        for cella, testo in zip(tabella.rows[0].cells, intestazioni):
            cella.paragraphs[0].add_run(testo).bold = True

        for voce in checklist:
            riga = tabella.add_row().cells
            riga[0].text = "☐"
            riga[1].text = voce.get("passo", "")
            riga[2].text = voce.get("responsabile", "")
            riga[3].text = "\n".join(
                f"- {d}" for d in voce.get("documenti", [])
            )
            tempi = voce.get("tempistica", "")
            note = voce.get("note", "")
            riga[4].text = "\n".join(parte for parte in (tempi, note) if parte)

    # Modulo / relazione
    modulo = dati.get("modulo", {})
    sezioni = modulo.get("sezioni", [])
    if sezioni:
        documento.add_page_break()
        p = documento.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(modulo.get("titolo", "Modulo"))
        run.bold = True
        run.font.size = Pt(14)
        p.paragraph_format.space_after = Pt(8)

        for sezione in sezioni:
            _titolo_sezione(documento, sezione.get("nome", ""))
            for campo in sezione.get("campi", []):
                _campo_compilabile(documento, campo)

        dichiarazioni = modulo.get("dichiarazioni", [])
        if dichiarazioni:
            _titolo_sezione(documento, "Dichiarazioni")
            for dich in dichiarazioni:
                documento.add_paragraph(dich, style="List Bullet")

        # Data e firme
        _titolo_sezione(documento, "Data e firme")
        _campo_compilabile(documento, "Luogo e data")
        documento.add_paragraph()
        for firma in modulo.get("firme", []):
            p = documento.add_paragraph()
            p.add_run(f"{firma}").bold = True
            p.add_run("\n" + "_" * 35)
            p.paragraph_format.space_after = Pt(10)

    documento.save(percorso_output)
    return percorso_output
